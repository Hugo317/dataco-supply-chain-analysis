"""Shared theme + slide-chrome helpers for the DataCo Supply Chain deck.

Palette matches the companion "Signal vs Noise" document (dark mode) so the
deck and the write-up read as a matched pair, distinct from the Vardane deck
this project's slide *organization* was modeled on.
"""

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette --
BG = "11161A"
SURFACE = "171E24"
SURFACE_2 = "1C242B"
INK = "E7ECEF"
INK_SOFT = "AAB4BC"
INK_FAINT = "6D7982"
LINE = "2A343C"
LINE_STRONG = "3A4650"
ACCENT = "6FA8D6"      # primary blue
AMBER = "E6A94A"       # secondary accent
CAUTION = "E08A5C"     # risk / conflict emphasis
GOOD = "5FBF90"        # positive / high confidence
CONF_COLORS = {"High": GOOD, "Medium": AMBER, "Low": INK_FAINT}

FONT = "Arial"

# ---------------------------------------------------------------- geometry -
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)
MARGIN_L = Emu(548640)
CONTENT_W = Emu(11091672)
EYEBROW_Y = Emu(347472)
HEADLINE_Y = Emu(621792)
RULE_Y = Emu(1691640)
BODY_TOP = Emu(1939694)
FOOTER_RULE_Y = Emu(6473952)
FOOTER_TEXT_Y = Emu(6537960)

CHART_X = Emu(548640)
CHART_Y = Emu(1939694)
CHART_W = Emu(7086600)   # 7.75in -- matches chart figsize=(7.75,4.26)
CHART_H = Emu(3893011)   # 4.26in
CARD_X = Emu(7882127)
CARD_Y = Emu(1920240)
CARD_W = Emu(3758184)
CARD_H = Emu(3931920)

BRAND = "DATACO SUPPLY CHAIN"


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    return slide


def add_rect(slide, x, y, w, h, fill=None, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    return shp


def add_round_rect(slide, x, y, w, h, fill, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    return shp


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=None, space_after=None, wrap=True):
    """runs: list of paragraphs; each paragraph is a list of (text, size_pt, bold, color_hex)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para_runs in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        if space_after is not None:
            p.space_after = Pt(space_after)
        for text, size, bold, color in para_runs:
            r = p.add_run()
            r.text = text
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = rgb(color)
    return box


def add_simple_text(slide, x, y, w, h, text, size, bold, color, align=PP_ALIGN.LEFT,
                     anchor=MSO_ANCHOR.TOP, line_spacing=None):
    return add_text(slide, x, y, w, h, [[(text, size, bold, color)]], align, anchor, line_spacing)


def add_eyebrow(slide, text, color=ACCENT):
    add_simple_text(slide, MARGIN_L, EYEBROW_Y, CONTENT_W, Emu(274320), text, 12, True, color)


def add_headline(slide, text, size=24, color=INK):
    add_simple_text(slide, MARGIN_L, HEADLINE_Y, CONTENT_W, Emu(1170432), text, size, True, color,
                     line_spacing=1.05)


def add_rule(slide, y=RULE_Y, color=LINE):
    add_rect(slide, MARGIN_L, y, CONTENT_W, Emu(9525), fill=color)


def add_accent_underline(slide, x, y):
    add_rect(slide, x, y, Emu(1188720), Emu(38100), fill=ACCENT)


def add_footer(slide, page_num):
    add_rect(slide, MARGIN_L, FOOTER_RULE_Y, CONTENT_W, Emu(9525), fill=LINE)
    add_simple_text(slide, MARGIN_L, FOOTER_TEXT_Y, Emu(3657600), Emu(274320), BRAND, 9, True, INK_FAINT)
    add_simple_text(slide, Emu(9811512), FOOTER_TEXT_Y, Emu(1828800), Emu(274320), str(page_num), 9, False,
                     INK_FAINT, align=PP_ALIGN.RIGHT)


def add_picture_fit(slide, path, x, y, w, h):
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def set_notes(slide, text):
    """No-op placeholder: this deck intentionally carries NO embedded speaker
    notes (a Keynote-compatibility constraint, same as the reference deck) --
    the talk track lives entirely in the companion document."""
    return
