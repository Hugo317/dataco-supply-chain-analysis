from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import pptx_theme as th


def title_slide(prs, kicker, title, subtitle, footnote):
    s = th.blank_slide(prs)
    th.add_simple_text(s, th.MARGIN_L, Emu(1200000), Emu(9601200), Emu(320040), kicker, 14, True, th.ACCENT)
    th.add_simple_text(s, th.MARGIN_L, Emu(1565000), Emu(10515600), Emu(1554480), title, 40, True, th.INK,
                        line_spacing=1.05)
    th.add_accent_underline(s, Emu(566928), Emu(3050000))
    th.add_simple_text(s, Emu(566928), Emu(3230000), Emu(10515600), Emu(600000), subtitle, 14, False, th.INK_SOFT)
    th.add_simple_text(s, th.MARGIN_L, Emu(6250000), Emu(9601200), Emu(365760), footnote, 11, False, th.INK_FAINT)
    return s


_NUMBER_WORDS = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven", 8: "Eight"}


def agenda_slide(prs, stories, page):
    s = th.blank_slide(prs)
    th.add_eyebrow(s, "AGENDA")
    n = _NUMBER_WORDS.get(len(stories), str(len(stories)))
    th.add_headline(s, f"{n} findings, {n.lower()} calls to make")
    th.add_rule(s)
    y = Emu(1874519)
    row_h = Emu(722376)
    for i, (num, title, question) in enumerate(stories):
        ry = Emu(y.emu + i * row_h.emu)
        th.add_rect(s, th.MARGIN_L, Emu(ry.emu + row_h.emu - Emu(9525).emu), th.CONTENT_W, Emu(9525), fill=th.LINE)
        th.add_simple_text(s, th.MARGIN_L, ry, Emu(685800), row_h, num, 20, True, th.ACCENT,
                            anchor=MSO_ANCHOR.MIDDLE)
        th.add_simple_text(s, Emu(1371600), ry, Emu(2834640), row_h, title, 15, True, th.INK,
                            anchor=MSO_ANCHOR.MIDDLE)
        th.add_simple_text(s, Emu(4434840), ry, Emu(7223760), row_h, question, 14, False, th.INK_SOFT,
                            anchor=MSO_ANCHOR.MIDDLE)
    th.add_footer(s, page)
    return s


def headline_slide(prs, headline_title, chart_path, tiles, framing, page):
    """Chart (e.g. sales/profit by year) on the left, up to 4 stat tiles stacked
    2x2 on the right -- mirrors the data_slide chart+card layout used everywhere
    else in the deck."""
    s = th.blank_slide(prs)
    th.add_eyebrow(s, "DATA & HEADLINE NUMBERS")
    th.add_headline(s, headline_title)
    th.add_rule(s)
    if chart_path:
        th.add_picture_fit(s, chart_path, th.CHART_X, th.CHART_Y, th.CHART_W, th.CHART_H)

    cols, rows = 2, 2
    gap = Emu(114300)
    tile_w = Emu((th.CARD_W.emu - gap.emu * (cols - 1)) // cols)
    tile_h = Emu((th.CARD_H.emu - gap.emu * (rows - 1)) // rows)
    for i, (num, label) in enumerate(tiles[:4]):
        col, row = i % cols, i // cols
        x = Emu(th.CARD_X.emu + col * (tile_w.emu + gap.emu))
        y = Emu(th.CARD_Y.emu + row * (tile_h.emu + gap.emu))
        th.add_round_rect(s, x, y, tile_w, tile_h, fill=th.SURFACE)
        stat_size = 26 if len(num) <= 7 else 20
        th.add_simple_text(s, Emu(x.emu + 160000), Emu(y.emu + 150000), Emu(tile_w.emu - 320000), Emu(500000),
                            num, stat_size, True, th.ACCENT)
        th.add_simple_text(s, Emu(x.emu + 160000), Emu(y.emu + 640000), Emu(tile_w.emu - 320000),
                            Emu(tile_h.emu - 640000 - 100000), label, 10.5, False, th.INK_SOFT, line_spacing=1.05)
    box_y = Emu(th.CHART_Y.emu + th.CHART_H.emu + 91440)
    box_h = Emu(6473952 - box_y.emu - 45720)
    th.add_round_rect(s, th.MARGIN_L, box_y, th.CHART_W, box_h, fill=th.SURFACE_2)
    th.add_simple_text(s, Emu(th.MARGIN_L.emu + 228600), box_y, Emu(th.CHART_W.emu - 457200), box_h,
                        framing, 12, False, th.INK_SOFT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.15)
    th.add_footer(s, page)
    return s


def divider_slide(prs, story_num, total, section_title, question, page):
    s = th.blank_slide(prs)
    th.add_simple_text(s, th.MARGIN_L, Emu(1554480), Emu(3657600), Emu(1097280), f"0{story_num}", 60, True,
                        th.LINE_STRONG)
    th.add_simple_text(s, Emu(566928), Emu(2650000), Emu(9601200), Emu(822960), section_title, 34, True, th.INK)
    th.add_accent_underline(s, Emu(594360), Emu(3380000))
    th.add_simple_text(s, Emu(566928), Emu(3560000), Emu(9601200), Emu(822960), question, 16, False, th.ACCENT,
                        line_spacing=1.1)
    th.add_simple_text(s, Emu(566928), Emu(4550000), Emu(9601200), Emu(320040), f"INSIGHT {story_num} OF {total}",
                        11, True, th.INK_FAINT)
    th.add_footer(s, page)
    return s


def data_slide(prs, story_num, total, section_title, headline, chart_path, stat_number, stat_label,
                finding_text, page, stat_color=None):
    s = th.blank_slide(prs)
    th.add_eyebrow(s, f"INSIGHT {story_num} OF {total} — {section_title.upper()}")
    th.add_headline(s, headline)
    th.add_rule(s)
    if chart_path:
        th.add_picture_fit(s, chart_path, th.CHART_X, th.CHART_Y, th.CHART_W, th.CHART_H)
    th.add_round_rect(s, th.CARD_X, th.CARD_Y, th.CARD_W, th.CARD_H, fill=th.SURFACE)
    pad = Emu(292608)
    inner_w = Emu(th.CARD_W.emu - pad.emu * 2)
    stat_size = 38 if len(stat_number) <= 8 else 30 if len(stat_number) <= 13 else 24
    th.add_simple_text(s, Emu(th.CARD_X.emu + pad.emu), Emu(th.CARD_Y.emu + 292608), inner_w, Emu(822960),
                        stat_number, stat_size, True, stat_color or th.ACCENT)
    th.add_simple_text(s, Emu(th.CARD_X.emu + pad.emu), Emu(th.CARD_Y.emu + 1050000), inner_w, Emu(548640),
                        stat_label, 12.5, True, th.INK_SOFT, line_spacing=1.05)
    th.add_simple_text(s, Emu(th.CARD_X.emu + pad.emu), Emu(th.CARD_Y.emu + 1650000), inner_w,
                        Emu(th.CARD_H.emu - 1650000 - pad.emu), finding_text, 11.5, False, th.INK_SOFT,
                        line_spacing=1.25)
    th.add_footer(s, page)
    return s


def data_slide_wide(prs, story_num, total, section_title, headline, chart_path, chart_h, finding_text, page):
    """Chart spans full content width; a finding strip sits below it. Used when a
    chart needs more horizontal room than the two-column layout allows."""
    s = th.blank_slide(prs)
    th.add_eyebrow(s, f"INSIGHT {story_num} OF {total} — {section_title.upper()}")
    th.add_headline(s, headline)
    th.add_rule(s)
    if chart_path:
        th.add_picture_fit(s, chart_path, th.MARGIN_L, th.BODY_TOP, th.CONTENT_W, chart_h)
    strip_y = Emu(th.BODY_TOP.emu + chart_h.emu + 91440)
    strip_h = Emu(6473952 - strip_y.emu - 45720)
    th.add_round_rect(s, th.MARGIN_L, strip_y, th.CONTENT_W, strip_h, fill=th.SURFACE)
    th.add_simple_text(s, Emu(th.MARGIN_L.emu + 228600), strip_y, Emu(th.CONTENT_W.emu - 457200), strip_h,
                        finding_text, 12.5, False, th.INK_SOFT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.2)
    th.add_footer(s, page)
    return s


def recommendation_slide(prs, story_num, total, section_title, lead_text, do_text, caveat_text, page):
    s = th.blank_slide(prs)
    th.add_eyebrow(s, f"INSIGHT {story_num} OF {total} — {section_title.upper()}")
    th.add_headline(s, "Recommendation")
    th.add_rule(s)
    runs = [[(lead_text, 17.5, False, th.INK_SOFT)], [("", 8, False, th.INK_SOFT)]]
    if do_text:
        runs.append([("Do:  ", 17.5, True, th.ACCENT), (do_text, 17.5, False, th.INK_SOFT)])
    if caveat_text:
        runs.append([("", 8, False, th.INK_SOFT)])
        runs.append([("Watch for:  ", 17.5, True, th.INK_FAINT), (caveat_text, 17.5, False, th.INK_SOFT)])
    th.add_text(s, th.MARGIN_L, Emu(1965960), Emu(10424160), Emu(4206240), runs, line_spacing=1.25, space_after=6)
    th.add_footer(s, page)
    return s


def summary_slide(prs, headline, rows, page):
    s = th.blank_slide(prs)
    th.add_eyebrow(s, "SUMMARY")
    th.add_headline(s, headline)
    th.add_rule(s)
    headers = ["Finding", "Recommended action", "Confidence"]
    col_x = [th.MARGIN_L.emu, 3337560, 9738360]
    col_w = [2697480, 6300000, 1810512]
    hy = Emu(1874519)
    for i, htext in enumerate(headers):
        th.add_simple_text(s, Emu(col_x[i]), hy, Emu(col_w[i]), Emu(320040), htext, 11.5, True, th.INK_FAINT)
    row_top = Emu(2221991)
    row_h = Emu(841248)
    gap = Emu(9525)
    for i, (finding, action, conf) in enumerate(rows):
        ry = Emu(row_top.emu + i * (row_h.emu + gap.emu))
        th.add_rect(s, th.MARGIN_L, Emu(ry.emu - gap.emu), th.CONTENT_W, gap, fill=th.LINE)
        th.add_simple_text(s, Emu(col_x[0]), ry, Emu(col_w[0]), row_h, finding, 14, True, th.INK,
                            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
        th.add_simple_text(s, Emu(col_x[1]), ry, Emu(col_w[1]), row_h, action, 13, False, th.INK_SOFT,
                            anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
        color = th.CONF_COLORS.get(conf, th.INK_FAINT)
        th.add_simple_text(s, Emu(col_x[2]), ry, Emu(col_w[2]), row_h, conf, 14, True, color,
                            anchor=MSO_ANCHOR.MIDDLE)
    th.add_footer(s, page)
    return s


def thankyou_slide(prs, closing_prompt):
    s = th.blank_slide(prs)
    th.add_simple_text(s, th.MARGIN_L, Emu(2468880), Emu(9601200), Emu(1005840), "Thank you", 46, True, th.INK)
    th.add_accent_underline(s, Emu(566928), Emu(3291840))
    th.add_simple_text(s, Emu(566928), Emu(3474720), Emu(9601200), Emu(457200), closing_prompt, 16, False,
                        th.ACCENT)
    return s


def appendix_divider(prs, subtitle, page):
    s = th.blank_slide(prs)
    th.add_simple_text(s, th.MARGIN_L, Emu(3017520), Emu(10515600), Emu(914400), "Appendix", 40, True, th.INK)
    th.add_simple_text(s, Emu(566928), Emu(3749039), Emu(10515600), Emu(457200), subtitle, 15, False, th.INK_SOFT)
    th.add_footer(s, page)
    return s


def appendix_text_slide(prs, headline, items, page):
    """items: list of (label, body) pairs."""
    s = th.blank_slide(prs)
    th.add_eyebrow(s, "APPENDIX")
    th.add_headline(s, headline)
    th.add_rule(s)
    runs = []
    for label, body in items:
        runs.append([(label + "  ", 13, True, th.INK), (body, 13, False, th.INK_SOFT)])
        runs.append([("", 8, False, th.INK_SOFT)])
    th.add_text(s, th.MARGIN_L, Emu(1828800), Emu(11064240), Emu(4400000), runs, line_spacing=1.2, space_after=4)
    th.add_footer(s, page)
    return s
